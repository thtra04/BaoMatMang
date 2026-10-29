"""Sinh slide thuyết trình (.pptx).   Chạy:  python slide/tao_slide.py"""
import ast
import os
from collections import Counter

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "so_do")
OUT = os.path.join(ROOT, "Slide_DoAn_Huong1_SafePayGuard.pptx")

NAVY = RGBColor(0x13, 0x23, 0x3A)
NAVY2 = RGBColor(0x22, 0x3A, 0x5E)
INK = RGBColor(0x1D, 0x24, 0x33)
MUTED = RGBColor(0x5B, 0x66, 0x7A)
LIGHT = RGBColor(0xF1, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xD7, 0x26, 0x3D)
REDBG = RGBColor(0xFC, 0xE9, 0xEB)
GREEN = RGBColor(0x14, 0x8F, 0x77)
GREENBG = RGBColor(0xE3, 0xF4, 0xF0)
AMBER = RGBColor(0xC9, 0x7C, 0x10)
AMBERBG = RGBColor(0xFD, 0xF3, 0xE1)
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def new_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY if dark else WHITE
    return s


def text(s, x, y, w, h, runs, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=0.05, space_after=6, bullets=False):
    """runs: chuỗi, hoặc list các đoạn; mỗi đoạn là chuỗi hoặc list[(text, {opts})]."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        parts = para if isinstance(para, list) else [(para, {})]
        if bullets:
            parts = [("•  ", {"color": RED})] + parts
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def shape(s, kind, x, y, w, h, fill, line=None, radius=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def badge(s, x, y, d, label, fill, color=WHITE, size=16):
    c = shape(s, MSO_SHAPE.OVAL, x, y, d, d, fill)
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = BODY
    r.font.color.rgb = color
    return c


def card(s, x, y, w, h, title, body, mark, mark_fill, fill=LIGHT, body_size=14, title_size=17, bullets=False, gap=4):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.08)
    badge(s, x + 0.25, y + 0.25, 0.55, mark, mark_fill, size=15)
    text(s, x + 0.95, y + 0.22, w - 1.15, 0.6, title, size=title_size, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 0.95, w - 0.5, h - 1.1, body, size=body_size, color=INK, space_after=gap, bullets=bullets)


def title(s, t, sub=None, dark=False):
    text(s, 0.6, 0.35, 12.1, 0.8, t, size=34, bold=True, font=HEAD, color=WHITE if dark else NAVY)
    if sub:
        text(s, 0.6, 1.12, 12.1, 0.45, sub, size=16, color=RGBColor(0xC8, 0xD3, 0xE3) if dark else MUTED)


def picture(s, fname, x, y, w, h):
    path = os.path.join(FIG, fname)
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    s.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def footer(s, n, dark=False):
    text(s, 0.6, 7.0, 8, 0.3, "SafePay Guard · Đồ án Bảo mật mạng máy tính và hệ thống", size=10,
         color=RGBColor(0x8A, 0x9B, 0xB5) if dark else MUTED)
    text(s, 11.9, 7.0, 0.8, 0.3, str(n), size=10, align=PP_ALIGN.RIGHT,
         color=RGBColor(0x8A, 0x9B, 0xB5) if dark else MUTED)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def table(s, x, y, w, col_w, header, rows, size=13, row_h=0.42, fills=None):
    shp = s.shapes.add_table(len(rows) + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(row_h * (len(rows) + 1)))
    t = shp.table
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for r in range(len(rows) + 1):
        t.rows[r].height = Inches(row_h)
        for c in range(len(header)):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = header[c] if r == 0 else rows[r - 1][c]
            cell.text = ""
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = val
            run.font.name = BODY
            run.font.size = Pt(size)
            run.font.bold = r == 0
            run.font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = NAVY
            else:
                cell.fill.fore_color.rgb = (fills(r - 1, c) if fills else None) or (LIGHT if r % 2 == 0 else WHITE)
    return t


n = 0

# ======================================================== 1. Bìa
s = new_slide(dark=True)
n += 1
badge(s, 0.9, 1.35, 0.9, "✓", GREEN, size=30)
text(s, 0.9, 2.45, 11, 1.1, "SafePay Guard", size=60, bold=True, font=HEAD, color=WHITE)
text(s, 0.9, 3.55, 11.2, 1.2, "Hệ thống xác minh và phê duyệt lệnh chuyển tiền chống giả mạo (deepfake, BEC) cho doanh nghiệp vừa và nhỏ",
     size=22, color=RGBColor(0xC8, 0xD3, 0xE3))
text(s, 0.9, 5.1, 11, 0.4, "Đồ án cuối kỳ · Bảo mật mạng máy tính và hệ thống · Hướng 1 – Security", size=16,
     color=RGBColor(0x8A, 0x9B, 0xB5))
text(s, 0.9, 5.55, 11, 0.4, "Nhóm: ……………………  ·  GVHD: ……………………", size=16, color=RGBColor(0x8A, 0x9B, 0xB5))
notes(s, "Giới thiệu nhóm và đề tài. Một câu tóm tắt: chúng em thiết kế hệ thống để một cuộc gọi deepfake của giám đốc, "
         "dù giống thật đến đâu, cũng không thể khiến tiền của công ty bị chuyển đi.")

# ======================================================== 2. Bài toán
s = new_slide()
n += 1
title(s, "Bài toán: tiền bị chuyển vì yêu cầu “trông có vẻ thật”")
for i, (num, lbl, src, col) in enumerate([
    ("2,9 tỷ USD", "thiệt hại do lừa đảo email doanh nghiệp (BEC) tại Mỹ năm 2023", "FBI IC3, Internet Crime Report 2023", RED),
    ("25 triệu USD", "một nhân viên Arup (Hồng Kông) chuyển đi sau cuộc họp video toàn người deepfake", "CNN, 02/2024", RED),
    ("Zalo · video", "thủ đoạn giả danh lãnh đạo yêu cầu kế toán chuyển tiền gấp tại Việt Nam", "Cảnh báo của cơ quan công an", AMBER),
]):
    x = 0.6 + i * 4.1
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.7, 3.8, 3.0, LIGHT, radius=0.06)
    text(s, x + 0.3, 1.95, 3.3, 0.9, num, size=38, bold=True, font=HEAD, color=col)
    text(s, x + 0.3, 2.9, 3.3, 1.2, lbl, size=15)
    text(s, x + 0.3, 4.15, 3.3, 0.4, src, size=11, color=MUTED)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 5.0, 12.1, 1.65, NAVY, radius=0.08)
text(s, 0.95, 5.15, 11.5, 1.4, [
    [("Điểm yếu chung: ", {"bold": True, "color": WHITE}),
     ("quy trình tin vào giọng nói, khuôn mặt, địa chỉ email, chữ ký trên văn bản – tất cả đều làm giả được với chi phí ngày càng thấp.",
      {"color": WHITE})],
    [("Câu hỏi của đồ án: ", {"bold": True, "color": RGBColor(0x8F, 0xE3, 0xCF)}),
     ("hệ thống dựa vào đâu để tin, kiểm chứng hoặc từ chối một yêu cầu chuyển tiền?", {"color": WHITE})],
], size=17, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, "Nêu bối cảnh bằng số liệu. Nhấn mạnh: vấn đề không nằm ở việc nhân viên kém cẩn thận, mà ở chỗ quy trình mặc định "
         "tin vào những thứ có thể làm giả.")

# ======================================================== 3. Người dùng & phạm vi
s = new_slide()
n += 1
title(s, "Ai dùng hệ thống, ai tấn công", "Doanh nghiệp 20–200 nhân sự, chuyển tiền cho nhà cung cấp qua Internet Banking hằng tuần")
users = [("Kế toán viên", "Lập lệnh, tiếp nhận yêu cầu đổi STK. Không được duyệt."),
         ("Kế toán trưởng", "Duyệt lệnh, gọi lại xác minh, duyệt đổi STK."),
         ("Giám đốc / Phó GĐ", "Bắt buộc với lệnh ≥ 500 triệu hoặc rủi ro cao."),
         ("Quản trị IT", "Duyệt thiết bị mới, khoá khẩn cấp. Không có quyền nghiệp vụ.")]
for i, (t, b) in enumerate(users):
    y = 1.85 + i * 1.22
    badge(s, 0.6, y, 0.6, str(i + 1), GREEN)
    text(s, 1.4, y - 0.05, 4.6, 0.4, t, size=17, bold=True, font=HEAD)
    text(s, 1.4, y + 0.36, 4.6, 0.7, b, size=14, color=MUTED)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.5, 1.75, 6.2, 4.95, REDBG, radius=0.05)
text(s, 6.85, 1.95, 5.6, 0.5, "Tác nhân đe dọa", size=20, bold=True, font=HEAD, color=RED)
text(s, 6.85, 2.55, 5.6, 4.0, [
    [("Nhóm lừa đảo BEC: ", {"bold": True}), ("email tên miền na ná, chiếm hộp thư NCC, giả giọng/video lãnh đạo", {})],
    [("Kẻ chiếm tài khoản: ", {"bold": True}), ("trang đăng nhập giả lấy mật khẩu + OTP, đánh cắp cookie", {})],
    [("Mã độc máy trạm: ", {"bold": True}), ("tráo số tiền / STK hiển thị trên trình duyệt", {})],
    [("Nội gián: ", {"bold": True}), ("tự lập – tự duyệt, sửa CSDL, xoá nhật ký", {})],
    [("Kẻ trộm thiết bị: ", {"bold": True}), ("lấy điện thoại của người duyệt", {})],
], size=15, bullets=True, space_after=9)
footer(s, n)
notes(s, "Giới thiệu 4 vai trò và nguyên tắc tách biệt nhiệm vụ. Bên phải là 5 nhóm tác nhân, ứng với các rủi ro ở slide sau.")

# ======================================================== 4. Tài sản & rủi ro
s = new_slide()
n += 1
title(s, "Tài sản cần bảo vệ và rủi ro chính")
assets = ["Tiền trong tài khoản doanh nghiệp", "Sổ đăng ký nhà cung cấp (STK, SĐT đã xác minh)",
          "Nội dung lệnh: người nhận, STK, số tiền", "Mật khẩu, bí mật OTP, khoá riêng thiết bị",
          "Khoá HMAC máy chủ, phiên đăng nhập", "Danh bạ nội bộ, nhật ký kiểm toán"]
picture(s, "05_ma_tran_rui_ro.png", 0.6, 1.2, 12.1, 4.45)
text(s, 0.6, 5.75, 1.6, 0.4, "Tài sản:", size=16, bold=True, font=HEAD, color=NAVY)
for i, a in enumerate(assets):
    x = 2.0 + (i % 3) * 3.6
    y = 5.75 + (i // 3) * 0.55
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.45, 0.45, LIGHT, radius=0.3)
    text(s, x + 0.12, y, 3.25, 0.45, a, size=12, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, "R1 giả danh lãnh đạo và R2 email giả NCC có điểm cao nhất (20) – đây là trọng tâm thiết kế. Các rủi ro về chiếm tài "
         "khoản và sửa dữ liệu được xử lý bằng các lớp phòng thủ bổ sung.")

# ======================================================== 5. Ý tưởng cốt lõi
s = new_slide(dark=True)
n += 1
text(s, 0.6, 0.5, 12.1, 1.0, "Không cố phát hiện deepfake – làm cho deepfake vô dụng", size=30, bold=True, font=HEAD, color=WHITE)
text(s, 0.6, 1.45, 12.1, 0.5, "Mọi thông tin từ bên ngoài chỉ là lời khẳng định chưa được xác minh", size=18,
     color=RGBColor(0xC8, 0xD3, 0xE3))
cols = [("HỆ THỐNG TIN", GREEN, RGBColor(0x1C, 0x3B, 0x45), ["Chữ ký Ed25519 từ thiết bị đã đăng ký, khớp đúng nội dung lệnh",
                                                         "Sổ đăng ký NCC & danh bạ nội bộ (lập khi ký hợp đồng / nhận việc)",
                                                         "Kết quả gọi lại tới số lấy từ sổ đăng ký",
                                                         "Dữ liệu lệnh – sau khi kiểm tra lại HMAC và chữ ký"]),
        ("HỆ THỐNG KHÔNG TIN", RED, RGBColor(0x3D, 0x22, 0x33), ["Giọng nói, khuôn mặt trong cuộc gọi thoại / video",
                                                              "Email, tin nhắn, văn bản scan, hoá đơn ảnh",
                                                              "STK / số điện thoại do người yêu cầu cung cấp",
                                                              "Lời tự xưng “tôi là Giám đốc”, dữ liệu trình duyệt gửi lên"])]
for i, (h, col, bg, items) in enumerate(cols):
    x = 0.6 + i * 6.2
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.3, 5.9, 4.3, bg, radius=0.05)
    badge(s, x + 0.35, 2.55, 0.65, "✓" if i == 0 else "✕", col, size=20)
    text(s, x + 1.2, 2.6, 4.5, 0.55, h, size=20, bold=True, font=HEAD, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.35, 3.45, 5.3, 3.0, items, size=16, color=WHITE, bullets=True, space_after=12)
footer(s, n, dark=True)
notes(s, "Đây là thông điệp quan trọng nhất. Phát hiện deepfake là cuộc đua vũ trang và không bao giờ đạt 100%. Thay vào đó, "
         "tiền chỉ được chuyển khi có bằng chứng mật mã từ đúng người – thứ mà deepfake không tạo ra được.")

# ======================================================== 6. Kiến trúc
s = new_slide()
n += 1
title(s, "Kiến trúc và ranh giới tin cậy")
picture(s, "01_kien_truc_he_thong.png", 0.5, 1.25, 12.3, 5.65)
footer(s, n)
notes(s, "Bốn vùng tin cậy. Dữ liệu từ vùng đỏ chỉ vào hệ thống qua tay người nhập, kèm nhãn kênh nguồn. Điện thoại người duyệt "
         "là kênh thứ hai độc lập với trình duyệt. Máy chủ vẫn không tin CSDL của chính nó: kiểm tra lại HMAC trước khi thực thi. "
         "Sổ đăng ký NCC là mỏ neo tin cậy.")

# ======================================================== 7. Luồng phê duyệt
s = new_slide()
n += 1
title(s, "Luồng tạo – duyệt – thực thi một lệnh")
picture(s, "02_luong_phe_duyet_lenh.png", 0.4, 1.25, 7.9, 5.65)
for i, (h, b, col) in enumerate([
    ("STK lấy từ sổ đăng ký", "Không bao giờ lấy từ tin nhắn. STK lạ trong tin nhắn → chặn ngay.", NAVY2),
    ("Ký trên điện thoại", "Người duyệt thấy đúng nội dung sẽ ký; chữ ký phủ cả nonce và quyết định.", AMBER),
    ("Kiểm tra lại khi thực thi", "Không tin trạng thái APPROVED trong CSDL: HMAC, mọi chữ ký, thiết bị, STK.", RED),
]):
    y = 1.5 + i * 1.75
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.6, y, 4.15, 1.55, LIGHT, radius=0.08)
    badge(s, 8.8, y + 0.2, 0.5, str(i + 1), col, size=14)
    text(s, 9.45, y + 0.15, 3.2, 0.5, h, size=16, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 8.8, y + 0.75, 3.85, 0.8, b, size=13.5, color=MUTED)
footer(s, n)
notes(s, "Đi theo các bước 1–14. Nhấn vào bước 4 (STK từ sổ), bước 6–8 (WYSIWYS trên kênh thứ hai) và bước 12 (kiểm tra lại).")

# ======================================================== 8. Cơ chế bảo vệ
s = new_slide()
n += 1
title(s, "Sáu cơ chế bảo vệ chính")
mechs = [
    ("MFA gắn thiết bị", "Argon2id + TOTP chống phát lại + thiết bị ký challenge. Thiết bị mới cần admin duyệt.", "1", NAVY2),
    ("Chữ ký giao dịch", "Ed25519 trên JSON chuẩn hoá: nội dung, nonce, hạn, quyết định. Không ký được → không duyệt được.", "2", NAVY2),
    ("Phê duyệt theo rủi ro", "≥ 50tr: 2 người; ≥ 500tr hoặc rủi ro cao: có GĐ. Người tự xưng phải tự ký.", "3", AMBER),
    ("Gọi lại ngoài băng", "Chỉ chấp nhận số trong danh bạ / sổ đăng ký. Bị phủ nhận → REJECTED + cảnh báo.", "4", AMBER),
    ("Sổ đăng ký + 48h chờ", "Không có ô nhập STK tự do. Đổi STK: gọi lại, ký duyệt, chờ 48h, +30 điểm 30 ngày.", "5", GREEN),
    ("Toàn vẹn & nhật ký", "HMAC trên lệnh (khoá trong HSM). Nhật ký chuỗi băm + HMAC + mốc neo bên ngoài.", "6", GREEN),
]
for i, (h, b, m, col) in enumerate(mechs):
    x = 0.6 + (i % 3) * 4.1
    y = 1.45 + (i // 3) * 2.75
    card(s, x, y, 3.85, 2.5, h, b, m, col, body_size=14)
footer(s, n)
notes(s, "Mỗi cơ chế giải quyết một nhóm rủi ro. 1–2 chống chiếm tài khoản và sửa nội dung; 3–4 chống giả danh lãnh đạo; "
         "5 chống BEC đổi tài khoản; 6 chống nội gián và che dấu vết.")

# ======================================================== 9. Chấm điểm rủi ro
s = new_slide()
n += 1
title(s, "Chấm điểm rủi ro: đầu vào bên ngoài chỉ làm tăng điểm")
rules = [["Yêu cầu từ kênh ngoài (email, chat, ảnh, thoại, video)", "+25"],
         ["Tự xưng là lãnh đạo qua kênh ngoài", "+20"],
         ["Khẩn cấp / bí mật (“gấp”, “đừng báo ai”)", "+15"],
         ["Lần đầu trả cho NCC / số tiền > 3× trung bình", "+20"],
         ["STK NCC đổi trong 30 ngày", "+30"],
         ["Ngoài giờ làm việc", "+10"]]
table(s, 0.6, 1.5, 6.6, [5.6, 1.0], ["Luật", "Điểm"], rules, size=14, row_h=0.6)
levels = [("THẤP < 30", "1 người duyệt (≥ 50 triệu: 2)", GREEN, GREENBG),
          ("TRUNG BÌNH 30–59", "2 người duyệt", AMBER, AMBERBG),
          ("CAO ≥ 60", "2 người + Giám đốc + gọi lại", RED, REDBG),
          ("CHẶN", "STK lạ trong tin nhắn · NCC trong 48h chờ", RED, REDBG)]
for i, (h, b, col, bg) in enumerate(levels):
    y = 1.5 + i * 1.2
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.6, y, 5.1, 1.0, bg, radius=0.12)
    text(s, 7.85, y + 0.1, 4.7, 0.4, h, size=16, bold=True, color=col)
    text(s, 7.85, y + 0.5, 4.7, 0.4, b, size=14)
text(s, 0.6, 6.25, 12.1, 0.5, "Ví dụ: “Giám đốc” gọi điện yêu cầu trả gấp NCC mới → 25 + 20 + 15 + 20 = 80 điểm → rủi ro CAO.",
     size=15, color=MUTED)
footer(s, n)
notes(s, "Bộ luật minh bạch, giải thích được cho kế toán và kiểm toán viên. Điểm quan trọng: kẻ gian không thể làm yêu cầu trông "
         "an toàn hơn bằng cách thêm thông tin, vì không luật nào trừ điểm.")

# ======================================================== 10. Tình huống deepfake
s = new_slide()
n += 1
title(s, "Tình huống 1: “Giám đốc” gọi video yêu cầu chuyển gấp",
      "Hình ảnh và giọng nói giống thật – hệ thống không cần biết đó là deepfake")
steps = [("Kế toán nhập lệnh 300 triệu", "kênh: cuộc gọi · tự xưng: Giám đốc Nam · khẩn cấp", NAVY2),
         ("Rủi ro CAO – 80 điểm", "cần 2 người + chính GĐ Nam ký + gọi lại", AMBER),
         ("KTT và Phó GĐ đều duyệt", "vẫn CHƯA ĐỦ – thiếu chữ ký của GĐ Nam", AMBER),
         ("Gọi lại số “GĐ” đọc", "bị từ chối: không phải số trong danh bạ", RED),
         ("Gọi số GĐ trong danh bạ", "GĐ: “Tôi không gọi” → REJECTED, cảnh báo", GREEN)]
for i, (h, b, col) in enumerate(steps):
    x = 0.6 + i * 2.5
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, 2.25, 3.0, LIGHT, radius=0.08)
    badge(s, x + 0.2, 2.2, 0.55, str(i + 1), col, size=15)
    text(s, x + 0.2, 2.95, 1.9, 0.9, h, size=15, bold=True, font=HEAD)
    text(s, x + 0.2, 3.85, 1.9, 1.1, b, size=13, color=MUTED)
    if i < 4:
        a = shape(s, MSO_SHAPE.RIGHT_ARROW, x + 2.27, 3.35, 0.21, 0.3, MUTED)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 5.35, 12.1, 1.3, NAVY, radius=0.1)
text(s, 0.9, 5.45, 11.6, 1.1, [
    [("Biến thể: ", {"bold": True, "color": RGBColor(0x8F, 0xE3, 0xCF)}),
     ("kẻ gian đọc STK mới (9704 2299 8811) → STK khác sổ đăng ký → lệnh bị CHẶN ngay khi tạo.", {"color": WHITE})],
    [("Nếu đó là Giám đốc thật: ", {"bold": True, "color": RGBColor(0x8F, 0xE3, 0xCF)}),
     ("chỉ tốn thêm một lần bấm ký trên điện thoại của Giám đốc.", {"color": WHITE})],
], size=16, anchor=MSO_ANCHOR.MIDDLE)
footer(s, n)
notes(s, "Kịch bản KB03, KB04 – test TC16 đến TC19. Có thể demo trực tiếp trên web: tạo lệnh với kênh Cuộc gọi thoại, "
         "tự xưng Giám đốc, rồi thử gọi lại số 0988 777 666.")

# ======================================================== 11. BEC
s = new_slide()
n += 1
title(s, "Tình huống 2: email giả nhà cung cấp đổi số tài khoản")
picture(s, "04_luong_doi_so_tai_khoan.png", 0.5, 1.3, 12.3, 3.6)
text(s, 0.6, 5.05, 12.1, 1.8, [
    [("Email từ minhphat-packaging.co: ", {"bold": True}),
     ("“từ tháng này chuyển vào TK ACB 88886666999, liên hệ 0912 345 678”", {"italic": True})],
    [("Hệ thống: ", {"bold": True}),
     ("STK cũ giữ nguyên · gọi số trong email → từ chối · gọi số trong sổ (028 3812 3456) → NCC phủ nhận → REJECTED, cảnh báo CRITICAL", {})],
    [("Nếu NCC đổi thật: ", {"bold": True}), ("vẫn đổi được, nhưng 48 giờ đầu mọi lệnh tới NCC đó bị chặn", {})],
], size=15, bullets=True, space_after=8)
footer(s, n)
notes(s, "Kịch bản KB05, KB06 – test TC05, TC20, TC21, TC22. Lệnh đã duyệt trước khi đổi STK cũng bị chặn khi thực thi.")

# ======================================================== 12. Chiếm tài khoản
s = new_slide()
n += 1
title(s, "Tình huống 3: tài khoản bị chiếm quyền", "Mỗi lớp được thiết kế với giả định lớp trước đã bị vượt qua")
picture(s, "06_phong_thu_nhieu_lop.png", 0.5, 1.7, 12.3, 5.2)
footer(s, n)
notes(s, "Đi từ trái sang phải: mật khẩu → OTP → thiết bị → phiên → điện thoại của một người duyệt. Ở lớp cuối, khoá khẩn cấp "
         "thu hồi thiết bị, huỷ phiên và huỷ cả các chữ ký đã ký nhưng chưa thực thi (TC15).")

# ======================================================== 13. Dữ liệu bị sửa
s = new_slide()
n += 1
title(s, "Tình huống 4: dữ liệu bị chỉnh sửa")
items = [("Mã độc tráo số tiền", "Màn hình hiển thị 4,8 triệu thay vì 48 triệu → chữ ký không khớp lệnh → từ chối, cảnh báo.", "TC23"),
         ("Nội gián sửa CSDL", "Đổi STK của lệnh đã duyệt → HMAC sai → lệnh BLOCKED khi thực thi.", "TC24"),
         ("Chữ ký dùng sai ngữ cảnh", "Chữ ký lệnh A dùng cho lệnh B; chữ ký “từ chối” dùng làm “duyệt” → từ chối.", "TC25–26"),
         ("Sửa / xoá nhật ký", "Sửa bản ghi, xoá bản ghi giữa, cắt đuôi → phát hiện bằng chuỗi băm, HMAC, mốc neo.", "TC27")]
for i, (h, b, tc) in enumerate(items):
    x = 0.6 + (i % 2) * 6.15
    y = 1.5 + (i // 2) * 2.65
    card(s, x, y, 5.9, 2.35, h, b, "!", RED, body_size=15)
    text(s, x + 4.4, y + 1.85, 1.3, 0.35, tc, size=12, color=MUTED, align=PP_ALIGN.RIGHT)
footer(s, n)
notes(s, "Demo trên web: nút 'Tấn công: mã độc tráo số tiền', nút 'Sửa trực tiếp trong CSDL', và trang Nhật ký với nút sửa/xoá bản ghi.")

# ======================================================== 13b. Bốn dạng đầu vào không đáng tin
s = new_slide()
n += 1
title(s, "Bốn dạng đầu vào không đáng tin", "Giả mạo · chỉnh sửa · đánh cắp · sử dụng sai ngữ cảnh – mỗi dạng có cơ chế riêng")
kinds = [["Giả mạo", "Deepfake video/giọng nói, email giả NCC, ảnh hoá đơn bị sửa STK, ảnh chụp màn hình “GĐ đã duyệt”",
          "Chỉ là lời khẳng định: STK lấy từ sổ, người tự xưng phải tự ký, gọi lại số danh bạ", "KB03–05, KB15"],
         ["Chỉnh sửa", "Tráo số tiền lúc ký, sửa CSDL sau khi duyệt, sửa/xoá nhật ký",
          "Chữ ký phủ toàn bộ nội dung, HMAC kiểm tra lại khi thực thi, chuỗi băm + mốc neo", "KB10, 11, 14"],
         ["Đánh cắp", "Mật khẩu, OTP, cookie phiên, điện thoại người duyệt",
          "MFA gắn thiết bị, phê duyệt cần khoá riêng, người duyệt thứ hai, khoá khẩn cấp", "KB07–09, KB13"],
         ["Sai ngữ cảnh", "Chữ ký của lệnh khác, chữ ký “từ chối”, OTP cũ, phê duyệt quá hạn, số gọi lại của kẻ gian",
          "Ràng buộc ngữ cảnh: request_id, nonce, quyết định, hạn trong nội dung ký; bộ đếm OTP", "TC09, 18, 22, 25, 26, 30"]]
kind_bg = [REDBG, AMBERBG, RGBColor(0xE6, 0xEE, 0xFA), GREENBG]
t = table(s, 0.6, 1.75, 12.1, [1.9, 4.2, 4.3, 1.7], ["Dạng", "Ví dụ", "Hệ thống xử lý", "Kịch bản / test"], kinds,
          size=14, row_h=1.08, fills=lambda r, c: kind_bg[r] if c == 0 else None)
t.rows[0].height = Inches(0.5)
for r in range(1, 5):
    t.cell(r, 0).text_frame.paragraphs[0].runs[0].font.bold = True
footer(s, n)
notes(s, "Đề bài yêu cầu ít nhất một tình huống đầu vào bị giả mạo, chỉnh sửa, đánh cắp hoặc sử dụng sai ngữ cảnh – nhóm xử lý cả "
         "bốn dạng. Nhấn mạnh hai ý mới: ảnh hoá đơn bị sửa STK vẫn bị chặn vì STK luôn lấy từ sổ đăng ký (KB15, TC33); và dữ liệu "
         "thật nhưng dùng sai ngữ cảnh – như chữ ký hợp lệ của lệnh A đem dùng cho lệnh B – bị từ chối vì mọi chữ ký đều gắn với "
         "đúng lệnh, đúng quyết định và đúng thời hạn (báo cáo mục 5.8).")

# ======================================================== 14. Prototype
s = new_slide()
n += 1
title(s, "Prototype: Python + Flask", "Chạy được trên máy cá nhân, kèm 15 kịch bản demo tự động")
text(s, 0.6, 1.75, 5.6, 0.45, "Thành phần", size=20, bold=True, font=HEAD, color=NAVY)
text(s, 0.6, 2.3, 5.6, 4.3, [
    [("core.py  ", {"bold": True}), ("chính sách, phê duyệt, thực thi", {})],
    [("crypto.py  ", {"bold": True}), ("Argon2id, TOTP, Ed25519, HMAC", {})],
    [("risk.py  ", {"bold": True}), ("bộ chấm điểm rủi ro", {})],
    [("audit.py  ", {"bold": True}), ("nhật ký chuỗi băm + mốc neo", {})],
    [("device.py  ", {"bold": True}), ("mô phỏng ứng dụng ký trên điện thoại", {})],
    [("app.py  ", {"bold": True}), ("giao diện web + nút mô phỏng tấn công", {})],
], size=15, bullets=True, space_after=8)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.6, 1.75, 6.1, 4.9, NAVY, radius=0.05)
text(s, 6.9, 1.95, 5.6, 0.45, "Chạy thử", size=20, bold=True, font=HEAD, color=WHITE)
text(s, 6.9, 2.55, 5.6, 2.0, ["pip install flask cryptography argon2-cffi pytest", "cd prototype",
                              "python demo_cli.py", "python -m pytest -v", "python app.py  →  127.0.0.1:5000"],
     size=13, font="Courier New", color=RGBColor(0xDB, 0xE7, 0xFF), space_after=4)
text(s, 6.9, 4.65, 5.6, 1.9, [
    "Trang đăng nhập hiển thị OTP của từng tài khoản demo",
    "Màn hình “điện thoại” hiển thị đúng nội dung sẽ ký",
    "Trang “Kịch bản tấn công” chạy cả 15 kịch bản",
], size=14, color=WHITE, bullets=True, space_after=6)
footer(s, n)
notes(s, "Chuyển sang demo trực tiếp ở bước này nếu có thời gian: KB01 luồng bình thường, KB04 giả giọng giám đốc, "
         "KB10 tráo nội dung, KB14 sửa nhật ký.")

# ======================================================== 15. Kiểm thử
src = open(os.path.join(ROOT, "prototype", "tests", "test_safepay.py"), encoding="utf-8").read()
groups = Counter()
for node in ast.parse(src).body:
    if isinstance(node, ast.FunctionDef) and node.name.startswith("test_tc"):
        groups[(ast.get_docstring(node) or "").split("|")[0].strip()] += 1
s = new_slide()
n += 1
title(s, "Kiểm thử: 33 test case, 38/38 lượt đạt")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.5, 3.9, 5.15, NAVY, radius=0.05)
text(s, 0.9, 1.8, 3.4, 1.2, "38/38", size=60, bold=True, font=HEAD, color=WHITE)
text(s, 0.9, 3.0, 3.4, 0.5, "lượt kiểm thử đạt (100%)", size=16, color=RGBColor(0xC8, 0xD3, 0xE3))
text(s, 0.9, 3.75, 3.4, 2.7, [
    f"{groups['Bình thường']} test luồng bình thường",
    f"{sum(groups.values()) - groups['Bình thường']} test tình huống rủi ro",
    "Đồng hồ giả lập: hết hạn, khoá tạm, 48h chờ",
    "Kiểm tra đúng mã lỗi chính sách",
], size=14, color=WHITE, bullets=True, space_after=7)
cd = CategoryChartData()
order = ["Bình thường", "Chiếm tài khoản", "Đầu vào giả mạo", "Sửa đổi dữ liệu", "Phân quyền"]
cd.categories = order
cd.add_series("Số test case", [groups[g] for g in order])
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(4.9), Inches(1.5), Inches(7.8), Inches(5.15), cd).chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Số test case theo nhóm"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
ch.font.size = Pt(14)
ch.font.name = BODY
plot = ch.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(14)
plot.data_labels.font.bold = True
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = NAVY2
for idx in range(1, len(order)):
    pt = ser.points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = RED
ch.value_axis.has_major_gridlines = False
ch.value_axis.visible = False
ch.category_axis.reverse_order = True
ch.category_axis.format.line.color.rgb = RGBColor(0xC3, 0xCA, 0xD6)
footer(s, n)
notes(s, "Bảng test case đầy đủ ở chương 7 của báo cáo. Kết quả lưu trong thư mục ket_qua_kiem_thu. Lưu ý trung thực: kiểm thử "
         "chứng minh hệ thống thực hiện đúng chính sách trong các tình huống đã nêu, không chứng minh an toàn trước mọi tấn công.")

# ======================================================== 16. Trade-off
s = new_slide()
n += 1
title(s, "Trade-off: an toàn có cái giá của nó")
rows = [["Duyệt nhiều lớp, gọi lại", "Chậm lệnh gấp, người dùng phản ứng", "Thành thói quen; nhân viên có lý do để từ chối “sếp”", "Chỉ kích hoạt theo rủi ro"],
        ["48 giờ chờ khi đổi STK", "Có thể trễ hạn thanh toán", "Chặn gần như toàn bộ BEC đổi tài khoản", "Giữ; báo NCC từ khi ký HĐ"],
        ["Chữ ký thiết bị", "Chi phí làm app, đào tạo", "Vận hành rẻ, an toàn bền vững", "App điện thoại; FIDO2 cho GĐ"],
        ["Luật vs AI phát hiện deepfake", "Luật đơn giản, đôi khi báo nhầm", "AI lỗi thời nhanh, khó giải thích", "Luật quyết định; AI chỉ cộng điểm"],
        ["Nhật ký đầy đủ vs riêng tư", "Điều tra dễ", "Kho dữ liệu dễ bị lạm dụng", "Thu thập tối thiểu, có hạn lưu"],
        ["Tích hợp ngân hàng", "Lớp duyệt riêng + đối soát", "API ngân hàng trực tiếp", "Triển khai 2 giai đoạn"]]
table(s, 0.6, 1.45, 12.1, [2.8, 2.9, 3.4, 3.0], ["Đánh đổi", "Ngắn hạn", "Dài hạn", "Quyết định"], rows, size=13, row_h=0.74,
      fills=lambda r, c: GREENBG if c == 3 else None)
footer(s, n)
notes(s, "Trade-off quan trọng nhất là an toàn – tiện lợi. Ngắn hạn: chậm và gây phản ứng. Dài hạn: thành thói quen, và một vụ BEC "
         "bị chặn (hàng trăm triệu đến hàng tỷ đồng) bù đắp chi phí chậm trễ nhiều năm.")

# ======================================================== 17. Hạn chế
s = new_slide()
n += 1
title(s, "Hạn chế và hướng phát triển")
card(s, 0.6, 1.5, 5.9, 5.1, "Hạn chế", [
    "Phụ thuộc chất lượng sổ đăng ký ban đầu",
    "Không chống được thông đồng của đủ số người duyệt",
    "Lệnh < 50 triệu chỉ cần 1 người duyệt",
    "Kết quả gọi lại do con người khai",
    "Prototype: lưu bộ nhớ, chưa HTTPS, HSM, app thật",
], "!", RED, body_size=17, bullets=True, gap=14)
card(s, 6.8, 1.5, 5.9, 5.1, "Hướng phát triển", [
    "App di động với passkey/FIDO2, Secure Enclave",
    "Xác nhận tự động trên app của người được nêu tên",
    "Đối chiếu hoá đơn điện tử, tên chủ tài khoản",
    "Học máy phát hiện bất thường (chỉ cộng điểm)",
    "Mốc neo nhật ký lên Object Lock / RFC 3161",
], "→", GREEN, body_size=17, bullets=True, gap=14)
footer(s, n)
notes(s, "Trình bày thẳng thắn các rủi ro còn lại – đặc biệt là sự phụ thuộc vào dữ liệu gốc và rủi ro thông đồng.")

# ======================================================== 18. Kết luận
s = new_slide(dark=True)
n += 1
text(s, 0.6, 0.6, 12.1, 0.9, "Kết luận", size=40, bold=True, font=HEAD, color=WHITE)
for i, (h, b) in enumerate([
    ("Chuyển gốc rễ của niềm tin", "từ thứ “trông có vẻ thật” sang chữ ký số, dữ liệu gốc đã xác minh và kênh xác minh độc lập"),
    ("Phòng thủ nhiều lớp", "giả mạo đầu vào, chiếm tài khoản, sửa dữ liệu đều có lớp chặn riêng và được kiểm thử"),
    ("Thực tế, triển khai được", "lệnh thường ngày vẫn nhanh; chi phí thấp; lộ trình 2 giai đoạn"),
]):
    y = 1.85 + i * 1.35
    badge(s, 0.7, y, 0.65, str(i + 1), GREEN, size=18)
    text(s, 1.6, y - 0.05, 10.9, 0.45, h, size=20, bold=True, font=HEAD, color=WHITE)
    text(s, 1.6, y + 0.42, 10.9, 0.6, b, size=16, color=RGBColor(0xC8, 0xD3, 0xE3))
text(s, 0.6, 6.0, 12.1, 0.7, "Cảm ơn thầy/cô và các bạn đã lắng nghe – Hỏi & đáp", size=22, bold=True, color=RGBColor(0x8F, 0xE3, 0xCF))
footer(s, n, dark=True)
notes(s, "Nhắc lại thông điệp chính: deepfake hoàn hảo đến đâu cũng không tạo được chữ ký từ điện thoại của Giám đốc.")

prs.save(OUT)
print("saved", OUT, "slides:", n)
