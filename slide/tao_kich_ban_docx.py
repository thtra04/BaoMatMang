"""Chuyển kịch bản thuyết trình (Markdown) sang Word.
Dùng:  python slide/tao_kich_ban_docx.py"""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "KichBan_ThuyetTrinh_5phut.md")
OUT = os.path.join(os.path.dirname(HERE), "KichBan_ThuyetTrinh_5phut.docx")

NAVY = RGBColor(0x13, 0x23, 0x3A)
BLUE = RGBColor(0x1F, 0x4E, 0x8C)
RED = RGBColor(0xD7, 0x26, 0x3D)
MUTED = RGBColor(0x6B, 0x75, 0x88)
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(13)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name in ("Heading 1", "Heading 2", "Title"):
    st = doc.styles[name]
    st.font.name = FONT
    rf = st.element.get_or_add_rPr().get_or_add_rFonts()
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(k), None)
    rf.set(qn("w:eastAsia"), FONT)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*\([^)]*\)\*|\*[^*]+\*|`[^`]+`)")


def runs(par, text, size=None):
    """Thêm text có **đậm**, *nghiêng*, *(chỉ dẫn sân khấu)* và `mã` vào đoạn."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
            r.font.color.rgb = NAVY
        elif part.startswith("*("):
            r = par.add_run("[" + part[2:-2] + "]")
            r.italic = True
            r.font.color.rgb = RED
        elif part.startswith("*"):
            r = par.add_run(part[1:-1])
            r.italic = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Consolas")
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)


def table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(3.4), Cm(2.6), Cm(10.6)]
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.width = widths[ci]
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            runs(p, val, size=12)
            if ri == 0:
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(c, "1F4E8C")
            elif "trọng tâm" in row[2] or ri % 2 == 0:
                shade(c, "FDF3E1" if "trọng tâm" in row[2] else "F1F4F8")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if ln.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(ln[2:].upper())
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(12)
    elif ln.startswith("## "):
        h = doc.add_heading(level=2)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(ln[3:])
        r.font.size = Pt(14)
        r.font.color.rgb = BLUE
    elif ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip("|").split("|")]
            if not all(set(c) <= {"-", ":"} for c in cells):
                rows.append(cells)
            i += 1
        table(rows)
        continue
    elif ln == "---":
        p = doc.add_paragraph()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "8A9BB5")):
            bottom.set(qn(k), v)
        pBdr.append(bottom)
        p._p.get_or_add_pPr().append(pBdr)
    elif re.match(r"^\d+\. ", ln):
        runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
    elif ln.startswith("- "):
        runs(doc.add_paragraph(style="List Bullet"), ln[2:])
    elif ln:
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        runs(p, ln)
    i += 1

# Chú thích quy ước ngay dưới tiêu đề
note = doc.paragraphs[1].insert_paragraph_before()
runs(note, "Quy ước: **chữ đậm** = ý bắt buộc phải nói; *(chữ đỏ trong ngoặc vuông)* = thao tác bấm slide, không đọc thành lời.", size=11.5)
note.runs[0].font.color.rgb = MUTED

doc.save(OUT)
print("saved", OUT)
